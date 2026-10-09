import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(run, name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def set_para_format(p, space_before=0, space_after=6, line_spacing=1.15, left_indent=0):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)
    else:
        p.paragraph_format.left_indent = None

def update_para(p, text, bold_prefix="", indent_tab=True, space_after=6, left_indent=0):
    p.text = "" # Clear runs
    set_para_format(p, space_before=0, space_after=space_after, line_spacing=1.15, left_indent=left_indent)
    
    if indent_tab:
        r_tab = p.add_run("\t")
        set_run_font(r_tab, "Times New Roman", 12)
        
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        set_run_font(r_bold, "Times New Roman", 12, bold=True)
        
    r_body = p.add_run(text)
    set_run_font(r_body, "Times New Roman", 12)

print("Setup completed.")
